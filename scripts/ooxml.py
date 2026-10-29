"""Native PPTX primitives extracted and adapted from the reviewed HotRevoke builder.

All paths and source decks are supplied by build_report.py. No application or
repository is opened when this module is imported.
"""
import copy, hashlib, io, json, math, posixpath, re, zipfile
from pathlib import Path
from lxml import etree as E
NS={"p":"http://schemas.openxmlformats.org/presentationml/2006/main","a":"http://schemas.openxmlformats.org/drawingml/2006/main","r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships","c":"http://schemas.openxmlformats.org/drawingml/2006/chart","pr":"http://schemas.openxmlformats.org/package/2006/relationships","ct":"http://schemas.openxmlformats.org/package/2006/content-types"}
REL=NS['r']+'/'
EMU=914400
COLORS=['245680','B36B00','009E73','A23838','7B6093']


def tag(prefix, name):
    return f"{{{NS[prefix]}}}{name}"

def sub(parent, prefix, local_name, **attrs):
    return E.SubElement(parent, tag(prefix, local_name), {k: str(v) for k, v in attrs.items()})

def xml(el):
    order = {name: i for i, name in enumerate([
        "ln", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill",
        "effectLst", "effectDag", "highlight", "uLnTx", "uLn", "uFillTx", "uFill",
        "latin", "ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst",
    ])}
    for props in el.iter():
        if props.tag in {tag("a", "rPr"), tag("a", "defRPr"), tag("a", "endParaRPr")}:
            children = list(props)
            for child in children:
                props.remove(child)
            for child in sorted(children, key=lambda child: order.get(E.QName(child).localname, 99)):
                props.append(child)
    return E.tostring(el, encoding="UTF-8", xml_declaration=True, standalone=True)

def move(sp, x, y, w, h):
    xf = sp.find("p:spPr/a:xfrm", NS)
    xf.find("a:off", NS).set("x", str(round(x * EMU)))
    xf.find("a:off", NS).set("y", str(round(y * EMU)))
    xf.find("a:ext", NS).set("cx", str(round(w * EMU)))
    xf.find("a:ext", NS).set("cy", str(round(h * EMU)))

def tx_body(parent, text, size=18, color="000000", center=False):
    body = sub(parent, "a", "txBody")
    sub(body, "a", "bodyPr", anchor="ctr", lIns=90000, rIns=70000, tIns=35000, bIns=35000)
    sub(body, "a", "lstStyle")
    p = sub(body, "a", "p")
    pp = sub(p, "a", "pPr", algn="ctr" if center else "l")
    sub(pp, "a", "buNone")
    sub(sub(pp, "a", "lnSpc"), "a", "spcPct", val=110000)
    r = sub(p, "a", "r")
    rp = sub(r, "a", "rPr", sz=int(size * 100))
    sub(sub(rp, "a", "solidFill"), "a", "srgbClr", val=color)
    sub(rp, "a", "latin", typeface="Times New Roman")
    sub(rp, "a", "ea", typeface="宋体")
    sub(r, "a", "t").text = str(text)
    return body

def graphic_frame(slide, name, x, y, w, h, ident=90):
    tree = slide.find("p:cSld/p:spTree", NS)
    frame = sub(tree, "p", "graphicFrame")
    nv = sub(frame, "p", "nvGraphicFramePr")
    sub(nv, "p", "cNvPr", id=ident, name=name)
    sub(nv, "p", "cNvGraphicFramePr")
    sub(nv, "p", "nvPr")
    xf = sub(frame, "p", "xfrm")
    sub(xf, "a", "off", x=round(x * EMU), y=round(y * EMU))
    sub(xf, "a", "ext", cx=round(w * EMU), cy=round(h * EMU))
    return sub(frame, "a", "graphic")

def native_table(slide, rows, widths, y=3.03, height=2.75, size=18):
    if not rows or any(len(row) != len(widths) for row in rows):
        raise ValueError("Native table must have a consistent column count")
    g = graphic_frame(slide, "Editable evidence table", 0.82, y, sum(widths), height)
    gd = sub(g, "a", "graphicData", uri=NS["a"] + "/table")
    table = sub(gd, "a", "tbl")
    sub(table, "a", "tblPr", firstRow="1", bandRow="0")
    grid = sub(table, "a", "tblGrid")
    for width in widths:
        sub(grid, "a", "gridCol", w=round(width * EMU))
    for i, row in enumerate(rows):
        tr = sub(table, "a", "tr", h=round(height * EMU / len(rows)))
        for text in row:
            tc = sub(tr, "a", "tc")
            tx_body(tc, text, size, "245680" if i == 0 else "000000")
            pr = sub(tc, "a", "tcPr", marL=90000, marR=70000, marT=35000, marB=35000, anchor="ctr")
            for edge in ["lnL", "lnR", "lnT", "lnB"]:
                ln = sub(pr, "a", edge, w=12700)
                sub(sub(ln, "a", "solidFill"), "a", "srgbClr", val="A9BCCB")
            sub(sub(pr, "a", "solidFill"), "a", "srgbClr", val="E9F0F5" if i == 0 else "FFFFFF")

def chart_text(parent, text, size=14):
    rich = sub(parent, "c", "rich")
    sub(rich, "a", "bodyPr")
    sub(rich, "a", "lstStyle")
    p = sub(rich, "a", "p")
    r = sub(p, "a", "r")
    rp = sub(r, "a", "rPr", sz=int(size * 100))
    sub(rp, "a", "latin", typeface="Times New Roman")
    sub(rp, "a", "ea", typeface="宋体")
    sub(sub(rp, "a", "solidFill"), "a", "srgbClr", val="000000")
    sub(r, "a", "t").text = text

def chart_txpr(parent, size=14):
    pr = sub(parent, "c", "txPr")
    sub(pr, "a", "bodyPr")
    sub(pr, "a", "lstStyle")
    p = sub(pr, "a", "p")
    pp = sub(p, "a", "pPr")
    d = sub(pp, "a", "defRPr", sz=int(size * 100))
    sub(d, "a", "latin", typeface="Times New Roman")
    sub(d, "a", "ea", typeface="宋体")
    sub(sub(d, "a", "solidFill"), "a", "srgbClr", val="000000")
    sub(p, "a", "endParaRPr", lang="zh-CN")

def cached_ref(parent, kind, formula, values):
    ref = sub(parent, "c", kind + "Ref")
    sub(ref, "c", "f").text = formula
    cache = sub(ref, "c", kind + "Cache")
    if kind == "num":
        sub(cache, "c", "formatCode").text = "0.00"
    sub(cache, "c", "ptCount", val=len(values))
    for i, v in enumerate(values):
        if v is not None:
            sub(sub(cache, "c", "pt", idx=i), "c", "v").text = str(v)

def native_chart(slide, parts, chart_no, data, geometry=None):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    categories = data["categories"]
    series = data["series"]
    if not categories or not series:
        raise ValueError("Reviewed charts require categories and series")
    for item in series:
        if len(item["values"]) != len(categories):
            raise ValueError("Chart categories and values do not align")
        if any(v is not None and (not isinstance(v, (int, float)) or not math.isfinite(v)) for v in item["values"]):
            raise ValueError("Chart values must be finite numbers or null")

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append([data["x_title"]] + [item["name"] for item in series])
    for i, cat in enumerate(categories):
        ws.append([str(cat)] + [item["values"][i] for item in series])
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 24
    for col in range(2, len(series) + 2):
        ws.column_dimensions[get_column_letter(col)].width = 22
    out = io.BytesIO()
    wb.save(out)
    # Excel serializes numeric cells to its own decimal representation. Cache
    # those exact stored values so an edit/reopen cannot disagree with the chart.
    with zipfile.ZipFile(io.BytesIO(out.getvalue())) as xlsx:
        sheet_xml = E.fromstring(xlsx.read("xl/worksheets/sheet1.xml"))
    sn = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    stored_numbers = {}
    for cell in sheet_xml.findall(".//s:c", sn):
        value = cell.find("s:v", sn)
        if value is not None:
            stored_numbers[cell.get("r")] = value.text
    book_name = f"combined_chart{chart_no}.xlsx"
    parts[f"ppt/embeddings/{book_name}"] = out.getvalue()

    cs = E.Element(tag("c", "chartSpace"), nsmap={k: NS[k] for k in ["c", "a", "r"]})
    sub(cs, "c", "date1904", val="0")
    sub(cs, "c", "lang", val="zh-CN")
    chart = sub(cs, "c", "chart")
    if data.get("title"):
        title = sub(chart, "c", "title")
        chart_text(sub(title, "c", "tx"), data["title"], 17)
        sub(title, "c", "layout")
        sub(title, "c", "overlay", val="0")
        sub(chart, "c", "autoTitleDeleted", val="0")
    else:
        sub(chart, "c", "autoTitleDeleted", val="1")
    plot = sub(chart, "c", "plotArea")
    sub(plot, "c", "layout")
    is_bar = data.get("kind", "line") == "bar"
    ct = sub(plot, "c", "barChart" if is_bar else "lineChart")
    if is_bar:
        sub(ct, "c", "barDir", val="col")
    sub(ct, "c", "grouping", val="clustered" if is_bar else "standard")
    sub(ct, "c", "varyColors", val="0")
    n = len(categories)
    for j, item in enumerate(series):
        ser = sub(ct, "c", "ser")
        sub(ser, "c", "idx", val=j)
        sub(ser, "c", "order", val=j)
        col = get_column_letter(j + 2)
        cached_ref(sub(ser, "c", "tx"), "str", f"Data!${col}$1", [item["name"]])
        sp = sub(ser, "c", "spPr")
        color = item.get("color", COLORS[j % len(COLORS)])
        if is_bar:
            sub(sub(sp, "a", "solidFill"), "a", "srgbClr", val=color)
            sub(sub(sp, "a", "ln"), "a", "noFill")
            for k, point_color in enumerate(data.get("category_colors", [])):
                dp = sub(ser, "c", "dPt")
                sub(dp, "c", "idx", val=k)
                dsp = sub(dp, "c", "spPr")
                sub(sub(dsp, "a", "solidFill"), "a", "srgbClr", val=point_color)
                sub(sub(dsp, "a", "ln"), "a", "noFill")
        else:
            line = sub(sp, "a", "ln", w=28575)
            sub(sub(line, "a", "solidFill"), "a", "srgbClr", val=color)
            if item.get("dash"):
                sub(line, "a", "prstDash", val=item["dash"])
            mk = sub(ser, "c", "marker")
            sub(mk, "c", "symbol", val=["circle", "square", "triangle"][j % 3])
            sub(mk, "c", "size", val="7")
            ms = sub(mk, "c", "spPr")
            sub(sub(ms, "a", "solidFill"), "a", "srgbClr", val=color)
            sub(sub(ms, "a", "ln"), "a", "solidFill").append(E.Element(tag("a", "srgbClr"), val=color))
        cached_ref(sub(ser, "c", "cat"), "str", f"Data!$A$2:$A${n + 1}", [str(v) for v in categories])
        cached_values = [stored_numbers.get(f"{col}{i + 2}") for i in range(n)]
        cached_ref(sub(ser, "c", "val"), "num", f"Data!${col}$2:${col}${n + 1}", cached_values)
        if not is_bar:
            sub(ser, "c", "smooth", val="0")
    if is_bar:
        sub(ct, "c", "gapWidth", val="120")
        sub(ct, "c", "overlap", val="0")
    sub(ct, "c", "axId", val="100")
    sub(ct, "c", "axId", val="200")
    for name, axis_id, cross_id, pos, label in [
        ("catAx", 100, 200, "b", data["x_title"]),
        ("valAx", 200, 100, "l", data["y_title"]),
    ]:
        axis = sub(plot, "c", name)
        sub(axis, "c", "axId", val=axis_id)
        scaling = sub(axis, "c", "scaling")
        sub(scaling, "c", "orientation", val="minMax")
        if name == "valAx":
            sub(scaling, "c", "min", val=data.get("y_min", 0))
        sub(axis, "c", "delete", val="0")
        sub(axis, "c", "axPos", val=pos)
        if name == "valAx":
            grid = sub(axis, "c", "majorGridlines")
            line = sub(sub(grid, "c", "spPr"), "a", "ln", w="6350")
            sub(sub(line, "a", "solidFill"), "a", "srgbClr", val="DEDEDE")
        title = sub(axis, "c", "title")
        chart_text(sub(title, "c", "tx"), label, 14)
        sub(title, "c", "layout")
        sub(title, "c", "overlay", val="0")
        if name == "valAx":
            sub(axis, "c", "numFmt", formatCode=data.get("number_format", "0.0"), sourceLinked="0")
        sub(axis, "c", "majorTickMark", val="none")
        sub(axis, "c", "minorTickMark", val="none")
        sub(axis, "c", "tickLblPos", val="nextTo")
        chart_txpr(axis)
        sub(axis, "c", "crossAx", val=cross_id)
        sub(axis, "c", "crosses", val="autoZero")
        if name == "catAx":
            sub(axis, "c", "auto", val="1")
            sub(axis, "c", "lblAlgn", val="ctr")
            sub(axis, "c", "lblOffset", val="100")
        else:
            sub(axis, "c", "crossBetween", val="between")
    if data.get("show_legend", True):
        legend = sub(chart, "c", "legend")
        sub(legend, "c", "legendPos", val="b")
        sub(legend, "c", "layout")
        sub(legend, "c", "overlay", val="0")
        chart_txpr(legend, 14)
    sub(chart, "c", "plotVisOnly", val="1")
    sub(chart, "c", "dispBlanksAs", val="gap")
    sub(chart, "c", "showDLblsOverMax", val="0")
    ext = sub(cs, "c", "externalData", **{f"{{{NS['r']}}}id": "rId1"})
    sub(ext, "c", "autoUpdate", val="0")
    parts[f"ppt/charts/chart{chart_no}.xml"] = xml(cs)
    rels = E.Element(tag("pr", "Relationships"), nsmap={None: NS["pr"]})
    sub(rels, "pr", "Relationship", Id="rId1", Type=REL + "package", Target=f"../embeddings/{book_name}")
    parts[f"ppt/charts/_rels/chart{chart_no}.xml.rels"] = xml(rels)
    x, y, w, h = geometry or (0.80, 2.93, 11.8, 3.14)
    g = graphic_frame(slide, f"Editable result chart {chart_no}", x, y, w, h, ident=90 + chart_no)
    gd = sub(g, "a", "graphicData", uri=NS["c"])
    sub(gd, "c", "chart", **{f"{{{NS['r']}}}id": f"rIdE16Chart{chart_no}"})

def read_package(p):
    with zipfile.ZipFile(p) as z:return {n:z.read(n) for n in z.namelist()}

def relpath(part):return posixpath.join(posixpath.dirname(part),'_rels',posixpath.basename(part)+'.rels')

def resolve(part,target):return posixpath.normpath(posixpath.join(posixpath.dirname(part),target))

def text_style(sp, paragraphs, size=20, color='333333', bold=False, center=False, font='微软雅黑', after=12):
    tx=sp.find('p:txBody',NS)
    for p in tx.findall('a:p',NS):tx.remove(p)
    bp=tx.find('a:bodyPr',NS)
    bp.set('anchor','t');bp.set('wrap','square')
    for k in ['lIns','rIns','tIns','bIns']:bp.set(k,'0')
    for child in list(bp):
        if E.QName(child).localname in ['normAutofit','spAutoFit','noAutofit']:bp.remove(child)
    sub(bp,'a','noAutofit')
    if isinstance(paragraphs,str):paragraphs=paragraphs.split('\n')
    for text in paragraphs:
        p=sub(tx,'a','p');pp=sub(p,'a','pPr',algn='ctr' if center else 'l',marL=0,indent=0)
        sub(pp,'a','buNone');sub(sub(pp,'a','lnSpc'),'a','spcPct',val=115000)
        sub(sub(pp,'a','spcAft'),'a','spcPts',val=round(after*100))
        run=sub(p,'a','r');rp=sub(run,'a','rPr',sz=round(size*100),b='1' if bold else '0',lang='zh-CN')
        sub(sub(rp,'a','solidFill'),'a','srgbClr',val=color)
        sub(rp,'a','latin',typeface='Times New Roman' if font=='宋体' else font)
        sub(rp,'a','ea',typeface=font);sub(rp,'a','cs',typeface=font)
        sub(run,'a','t').text=text
    return sp

def textbox(slide,text,x,y,w,h,size=20,color='333333',bold=False,center=False,font='宋体',name='Weekly text'):
    tree=slide.find('p:cSld/p:spTree',NS)
    ids=[int(p.get('id')) for p in tree.findall('.//p:cNvPr',NS)]
    sp=sub(tree,'p','sp');nv=sub(sp,'p','nvSpPr')
    sub(nv,'p','cNvPr',id=max(ids+[1])+1,name=name);sub(nv,'p','cNvSpPr',txBox='1');sub(nv,'p','nvPr')
    props=sub(sp,'p','spPr');xf=sub(props,'a','xfrm')
    sub(xf,'a','off',x=round(x*EMU),y=round(y*EMU));sub(xf,'a','ext',cx=round(w*EMU),cy=round(h*EMU))
    sub(sub(props,'a','prstGeom',prst='rect'),'a','avLst');sub(props,'a','noFill');sub(sub(props,'a','ln'),'a','noFill')
    tx=sub(sp,'p','txBody');sub(tx,'a','bodyPr');sub(tx,'a','lstStyle')
    return text_style(sp,text,size,color,bold,center,font,0)
